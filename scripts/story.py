#!/usr/bin/env python3
"""故事线合集：把多个功能演示片串成一条片子，功能之间插入「故事卡」（开场 / 章节 / 结尾），
视觉取自用户提供的 PPT（底图、logo、字体），与演讲 / 发布会同源。用法与原则见 references/story-cards.md。

子命令（都在 <work>/_tmp 下、非沙箱运行；<work> = <ROOT>/<项目名>/work）：
  ppt      <deck.pptx> --out <项目>/assets/ppt   拆 PPT：底图/logo 联系表、版式→页面映射、字体排行、全文+备注大纲
  chapters <work> [key ...]                       按 story.json.chapters 用子项目 edit.json 重渲章节素材（无字幕、可裁切）
  plan     <work>                                 打印时间线、字幕表、总时长、<3s / >20 字告警
  cards    <work>                                 故事卡末帧静图 → <work>/checks/card_<key>.jpg
  build    <work> [--out x.mp4]                   出 <out_dir>/<name>_vN.mp4 + _vN.srt + _vN_无字幕.mp4

<work>/story.json 示例见 assets/story_example.json。
"""
import argparse, json, re, subprocess, sys, zipfile
from collections import Counter, deque
from pathlib import Path

import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
FPS = 30
MSYH = "/Applications/Microsoft Word.app/Contents/Resources/DFonts/msyh.ttc"
MSYHBD = "/Applications/Microsoft Word.app/Contents/Resources/DFonts/msyhbd.ttc"


# ================================================================== ppt：拆参考 PPT
def cmd_ppt(a):
    deck, out = Path(a.pptx).expanduser(), Path(a.out).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    z = zipfile.ZipFile(deck)
    names = z.namelist()

    def rels(part):
        p = re.sub(r"([^/]+)$", r"_rels/\1.rels", part)
        if p not in names:
            return {}
        x = z.read(p).decode("utf8", "ignore")
        return dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', x))

    # 版式 → 图片；页面 → 版式
    lay_media, slide_lay = {}, {}
    for n in names:
        if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n):
            lay_media[n] = [Path(t).name for t in rels(n).values() if "media/" in t]
    slides = sorted([n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)],
                    key=lambda s: int(re.findall(r"\d+", s)[-1]))
    for s in slides:
        for t in rels(s).values():
            if "slideLayout" in t:
                slide_lay[s] = "ppt/slideLayouts/" + Path(t).name
    # 导出版式/母版图片（版式里的大图通常就是封面/内容页底图）
    media = [n for n in names if re.search(r"ppt/(slideLayouts|slideMasters)/media/", n)
             or (n.startswith("ppt/media/"))]
    saved = []
    for n in media:
        if not re.search(r"\.(png|jpe?g)$", n, re.I):
            continue
        dst = out / Path(n).name
        dst.write_bytes(z.read(n))
        try:
            im = Image.open(dst)
            saved.append((dst, im.size, im.mode))
        except Exception:
            pass
    # 联系表
    if saved:
        tw, th = 480, 300
        cols = 4
        sheet = Image.new("RGB", (tw * cols, th * ((len(saved) + cols - 1) // cols)), (128, 128, 128))
        d = ImageDraw.Draw(sheet)
        for i, (p, sz, mode) in enumerate(saved):
            im = Image.open(p).convert("RGBA")
            im.thumbnail((tw - 8, th - 40))
            bg = Image.new("RGBA", im.size, (90, 90, 90, 255))   # 灰底看得出透明 logo
            bg.alpha_composite(im)
            x, y = (i % cols) * tw, (i // cols) * th
            sheet.paste(bg.convert("RGB"), (x + 4, y + 4))
            lf = font(MSYHBD, 20) if Path(MSYHBD).exists() else None
            d.text((x + 6, y + th - 34), f"{p.name} {sz[0]}x{sz[1]} {mode}", fill=(255, 255, 0), font=lf)
        sheet.save(out / "_media_sheet.jpg", quality=88)
    # 字体排行（按页面里实际使用次数）
    fonts = Counter()
    for s in slides:
        fonts.update(re.findall(r'typeface="([^"+][^"]*)"', z.read(s).decode("utf8", "ignore")))
    local = subprocess.run(["fc-list", ":", "family", "file"], capture_output=True, text=True).stdout
    # 全文 + 备注
    lines = [f"# {deck.name} 大纲\n"]
    try:
        from pptx import Presentation
        prs = Presentation(str(deck))
        for i, s in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text.strip().replace("\n", " / ") for sh in s.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip()]
            lines.append(f"## P{i}  （版式 {Path(slide_lay.get(slides[i - 1], '?')).stem}）")
            lines += [f"- {t}" for t in texts]
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                lines.append("> 备注：" + s.notes_slide.notes_text_frame.text.strip()[:800].replace("\n", " "))
            lines.append("")
    except Exception as e:
        lines.append(f"（python-pptx 读取失败：{e}；pip install python-pptx）")
    (out / "_outline.md").write_text("\n".join(lines), encoding="utf-8")
    # 汇总
    use = Counter(slide_lay.values())
    print(f"[ok] {len(saved)} 张图片 → {out}（联系表 _media_sheet.jpg，用 ViewImage 看）")
    print(f"[ok] 全文与备注 → {out / '_outline.md'}")
    print("[i] 版式使用次数（第 1 页的版式≈封面，用得最多的≈内容页）：")
    for lay, n in use.most_common():
        first = min(i for i, s in enumerate(slides, 1) if slide_lay.get(s) == lay)
        print(f"    {Path(lay).stem:16s} ×{n:<3d} 首次 P{first:<3d} 图片 {lay_media.get(lay, [])}")
    print("[i] 字体排行（页面实际使用次数 / 本机是否有）：")
    for f, n in fonts.most_common(8):
        hit = [l.split(":")[0] for l in local.splitlines() if f in l][:1]
        print(f"    {f:20s} ×{n:<4d} {'本机 ' + hit[0] if hit else '本机没有'}")


# ================================================================== 配置
def load_story(work):
    p = Path(work) / "story.json"
    if not p.exists():
        sys.exit(f"[x] 缺 {p}（模板见 {HERE.parent / 'assets/story_example.json'}）")
    st = json.loads(p.read_text(encoding="utf-8"))
    st["_work"] = Path(work).resolve()
    st["_proj"] = st["_work"].parent
    st.setdefault("size", [1920, 1080])
    f = st.setdefault("fonts", {})
    f.setdefault("title", MSYHBD)
    f.setdefault("text", MSYH)
    f.setdefault("caption", MSYHBD)
    for k in [k for k in f if k.startswith("_")]:
        f.pop(k)
    for k in ("chapters", "cards"):          # 「_说明」之类的注释字段
        st[k] = {kk: vv for kk, vv in (st.get(k) or {}).items() if not kk.startswith("_")}
    for k, v in f.items():
        f[k] = str(Path(v).expanduser())
        if not Path(f[k]).exists():
            sys.exit(f"[x] 字体不存在：fonts.{k} = {f[k]}")
    return st


def proj_path(st, p):
    p = Path(p).expanduser()
    return p if p.is_absolute() else st["_proj"] / p


# ================================================================== chapters：重渲章节素材
def cmd_chapters(a):
    """子项目已定稿时，章节素材必须从子项目 edit.json 重渲（裁切在推近之前做），不要对成片再裁。"""
    st = load_story(a.work)
    from from_video import PLAIN, make_keep, render_events, save_json, load_json, probe
    chs = st.get("chapters") or {}
    keys = a.keys or list(chs)
    od = st["_work"] / "chapters"
    od.mkdir(parents=True, exist_ok=True)
    for k in keys:
        c = chs[k]
        ed = load_json(proj_path(st, c["project"]) / "work" / "edit.json")
        video = Path(ed["video"])
        info = probe(video)
        steps = []
        for i, f in enumerate(ed.get("focus") or [], 1):
            if "box" in f:
                steps.append({"id": f"f{i}", "t0": f["t0"], "t1": f["t1"], "box": f["box"],
                              "zoom": f.get("zoom", 1.5), "caption": ""})
            else:
                steps.append({"id": f"f{i}", "t0": round(f["t"] - 0.8, 3), "t1": round(f["t"] + f.get("hold", 1.2), 3),
                              "point": {"t": f["t"], "x": f["x"], "y": f["y"], "kind": "click"},
                              "zoom": f.get("zoom", 1.8), "hold_zoom": True, "caption": ""})
        steps.sort(key=lambda s: s["t0"])
        style = dict(PLAIN)
        style.update(ed.get("style") or {})
        style["keep"] = make_keep(info["duration"], ed.get("cut"), ed.get("speed"))
        if ed.get("blur"):
            style["blur"] = ed["blur"]
        crop = c.get("crop", st.get("crop"))
        if crop:
            style["crop"] = crop
        ev = {"kind": "video", "video": str(video), "src_w": info["w"], "src_h": info["h"],
              "duration": info["duration"], "cursor_mode": "real", "steps": steps, "style": style}
        evp = od / f"{k}_events.json"
        save_json(evp, ev)
        render_events(evp, od / f"{k}.mp4")
        print(f"[ok] 章节 {k} → {od / (k + '.mp4')}")


# ================================================================== 绘制工具
_fc = {}


def font(path, size):
    if (path, size) not in _fc:
        _fc[(path, size)] = ImageFont.truetype(path, size)
    return _fc[(path, size)]


def ease(x):
    x = max(0.0, min(1.0, x))
    return 1 - (1 - x) ** 3


def text_layer(txt, path, size, fill=(255, 255, 255)):
    f = font(path, size)
    _, _, _, b = f.getbbox(txt)
    im = Image.new("RGBA", (int(f.getlength(txt)) + 4, b + 8), (0, 0, 0, 0))
    ImageDraw.Draw(im).text((0, 0), txt, font=f, fill=fill)
    return im


def gradient_text(txt, path, size, colors):
    m = np.array(text_layer(txt, path, size))
    h, w = m.shape[:2]
    xs = np.linspace(0, 1, w)
    g = np.zeros((h, w, 3), np.float32)
    for i in range(3):
        g[:, :, i] = np.interp(xs, np.linspace(0, 1, len(colors)), [c[i] for c in colors])[None, :]
    return Image.fromarray(np.dstack([g.astype(np.uint8), m[:, :, 3]]), "RGBA")


def pill(txt, path, size, fg=(255, 255, 255), bg=(255, 255, 255, 34), border=(255, 255, 255, 90), padx=26, pady=12):
    f = font(path, size)
    tw = int(f.getlength(txt))
    _, it, _, ib = f.getbbox("国")
    h = (ib - it) + pady * 2
    im = Image.new("RGBA", (tw + padx * 2, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((0, 0, im.width - 1, h - 1), radius=h // 2, fill=bg, outline=border, width=2)
    d.text((padx, pady - it), txt, font=f, fill=fg)
    return im


def blend_rgba(canvas, rgba, x, y, a):
    if a <= 0:
        return
    H, W = canvas.shape[:2]
    x, y = int(round(x)), int(round(y))
    h, w = rgba.shape[:2]
    x0, y0, x1, y1 = max(0, x), max(0, y), min(W, x + w), min(H, y + h)
    if x0 >= x1 or y0 >= y1:
        return
    src = rgba[y0 - y:y1 - y, x0 - x:x1 - x].astype(np.float32)
    al = src[:, :, 3:4] / 255.0 * a
    dst = canvas[y0:y1, x0:x1].astype(np.float32)
    canvas[y0:y1, x0:x1] = (dst * (1 - al) + src[:, :, :3] * al).astype(np.uint8)


def hexrgb(c):
    if isinstance(c, (list, tuple)):
        return tuple(c)
    c = c.lstrip("#")
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


# ================================================================== 故事卡
class Card:
    """layout: cover（开场）| chapter（章节）| end（结尾）。坐标按 1920 宽设计，其他尺寸等比缩放。"""

    def __init__(self, st, cfg):
        self.st, self.c = st, cfg
        self.W, self.H = st["size"]
        self.k = self.W / 1920
        self.n = int(round(cfg["dur"] * FPS))
        th = st.get("card_theme", {})
        self.grad = [hexrgb(c) for c in th.get("gradient", ["#C65CFF", "#4A6EFF", "#28D2E6"])]
        self.kicker_c = hexrgb(th.get("kicker", "#AABEFF"))
        self.sub_c = hexrgb(th.get("sub", "#D6DCEE"))
        self.accent = hexrgb(th.get("accent", "#4F6EFF"))
        bg = Image.open(proj_path(st, cfg["bg"])).convert("RGB")
        s = max(self.W / bg.width, self.H / bg.height) * 1.06          # 留 6% 给缓慢推近
        self.bg = np.array(bg.resize((int(bg.width * s + .5), int(bg.height * s + .5)), Image.LANCZOS))
        self.layers = []
        getattr(self, "build_" + cfg["layout"])()

    def S(self, v):
        return int(round(v * self.k))

    def add(self, im, x, y, t, slide=True):
        self.layers.append((np.array(im), self.S(x), self.S(y), t, slide))

    def logo(self):
        if self.c.get("logo"):
            lg = Image.open(proj_path(self.st, self.c["logo"])).convert("RGBA")
            h = self.S(self.c.get("logo_h", 96))
            lg = lg.resize((int(lg.width * h / lg.height), h), Image.LANCZOS)
            self.add(lg, 120, 84, 0.0, False)

    def build_cover(self):
        c, F = self.c, self.st["fonts"]
        self.logo()
        y = 300
        if c.get("kicker"):
            self.add(text_layer(c["kicker"], F["text"], self.S(40), self.kicker_c), 120, y, 0.15); y += 92
        for i, line in enumerate(c.get("title", [])):
            self.add(text_layer(line, F["title"], self.S(100)), 120, y, 0.35 + 0.15 * i); y += 132
        y += 40
        if c.get("sub"):
            # 副标题过长时自动缩小字号直到一行放得下（右侧留出与左侧相同的边距）
            sz, maxw = 36, self.W - self.S(120) * 2
            while sz > 24 and font(F["text"], self.S(sz)).getlength(c["sub"]) > maxw:
                sz -= 1
            self.add(text_layer(c["sub"], F["text"], self.S(sz), self.sub_c), 120, y, 0.8)
        x, t = 120, 1.1
        py0 = 905                                       # 标签胶囊顶边（1920 基准）
        for i, s in enumerate(c.get("flow", [])):
            p = pill(f"{i + 1:02d}  {s}", F["title"], self.S(32))
            self.add(p, x, py0, t)
            x += p.width / self.k + 18
            if i < len(c["flow"]) - 1:
                # 箭头按「实际渲染出的墨迹像素」与胶囊垂直居中（字体给的字框数据不可靠：箭头字形在框里偏下）
                ar = text_layer("→", F["text"], self.S(34), (200, 210, 240))
                al = np.array(ar)[:, :, 3]
                rows = np.where(al > 40)[0]
                ink_c = (rows.min() + rows.max()) / 2.0               # 箭头墨迹中心（箭头图层内，实际像素）
                ay = py0 + (p.height / 2.0 - ink_c) / self.k           # 胶囊中心 − 墨迹中心 → 图层顶边（1920 基准）
                self.add(ar, x, ay, t + 0.08)
                x += ar.width / self.k + 18
            t += 0.15

    def build_chapter(self):
        c, F = self.c, self.st["fonts"]
        if c.get("no"):
            self.add(gradient_text(c["no"], F["title"], self.S(150), self.grad), 150, 88, 0.0)
        self.add(text_layer(c["title"], F["title"], self.S(88)), 150, 262, 0.15)
        if c.get("story"):
            self.add(text_layer(c["story"], F["text"], self.S(44), self.sub_c), 150, 392, 0.4)
        if c.get("tag"):
            self.add(pill(c["tag"], F["title"], self.S(30), bg=self.accent + (120,),
                          border=(140, 165, 255, 200)), 150, 476, 0.65)

    def build_end(self):
        c, F = self.c, self.st["fonts"]
        self.logo()
        y = 290
        if c.get("kicker"):
            self.add(text_layer(c["kicker"], F["text"], self.S(40), self.kicker_c), 120, y, 0.15); y += 92
        for line in c.get("title", []):
            self.add(text_layer(line, F["title"], self.S(112)), 120, y, 0.3); y += 150
        if c.get("sub"):
            self.add(gradient_text(c["sub"], F["title"], self.S(56), self.grad), 120, y, 0.55)
        x = 120
        for i, (a, b) in enumerate(c.get("pillars", [])):
            w, h = self.S(380), self.S(132)
            im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
            d = ImageDraw.Draw(im)
            d.rounded_rectangle((0, 0, w - 1, h - 1), radius=self.S(18), fill=(10, 14, 40, 170),
                                outline=(150, 165, 255, 150), width=2)
            d.text((self.S(30), self.S(22)), a, font=font(F["title"], self.S(40)), fill=(255, 255, 255))
            d.text((self.S(30), self.S(78)), b, font=font(F["text"], self.S(30)), fill=(190, 200, 235))
            self.add(im, x, 820, 0.9 + 0.15 * i)
            x += 404

    def frame(self, i):
        bh, bw = self.bg.shape[:2]
        p = i / max(1, self.n - 1)
        cw = min(bw, bh * self.W / self.H) / (1.0 + 0.05 * p)        # 背景缓慢推近
        chh = cw * self.H / self.W
        x0, y0 = (bw - cw) / 2, (bh - chh) / 2
        M = np.float32([[self.W / cw, 0, -x0 * self.W / cw], [0, self.H / chh, -y0 * self.H / chh]])
        fr = cv2.warpAffine(self.bg, M, (self.W, self.H), flags=cv2.INTER_LINEAR)
        t = i / FPS
        for arr, x, y, t0, slide in self.layers:
            a = ease((t - t0) / 0.45)
            if a > 0:
                blend_rgba(fr, arr, x, y + ((1 - a) * self.S(26) if slide else 0), a)
        return fr

    def frames(self):
        for i in range(self.n):
            yield self.frame(i)


# ================================================================== 片段 / 时间线
def clip_len(c):
    return int(round((c[2] - c[1]) / c[3] * FPS))


def clip_frames(st, c):
    key, t0, t1, sp = c
    W, H = st["size"]
    n = clip_len(c)
    src = st["_work"] / "chapters" / f"{key}.mp4"
    if not src.exists():
        src = proj_path(st, (st.get("chapters", {}).get(key) or {}).get("file", src))
    vf = f"setpts=(PTS-STARTPTS)/{sp},fps={FPS},scale={W}:{H}:force_original_aspect_ratio=decrease:flags=lanczos," \
         f"pad={W}:{H}:(ow-iw)/2:(oh-ih)/2,setsar=1"
    p = subprocess.Popen(["ffmpeg", "-v", "quiet", "-ss", f"{t0:.3f}", "-i", str(src), "-t", f"{t1 - t0 + 0.5:.3f}",
                          "-vf", vf, "-an", "-f", "rawvideo", "-pix_fmt", "rgb24", "-"], stdout=subprocess.PIPE)
    sz, last = W * H * 3, None
    for _ in range(n):
        buf = p.stdout.read(sz)
        if len(buf) == sz:
            last = np.frombuffer(buf, np.uint8).reshape(H, W, 3)
        yield last                                   # 尾部不足一帧：重复最后一帧
    p.stdout.close()
    p.kill()


def plan(st):
    items, t = [], 0.0
    for k, it in enumerate(st["timeline"]):
        xf = it.get("xf", 0.0) if k else 0.0
        n = int(round(st["cards"][it["card"]]["dur"] * FPS)) if "card" in it else clip_len(it["clip"])
        start = t - xf
        items.append(dict(it, n=n, xfn=int(round(xf * FPS)), start=start))
        t = start + n / FPS
    caps = []
    for k, it in enumerate(items):
        if "clip" not in it:
            continue
        _, c0, _, sp = it["clip"]
        # 字幕不压在故事卡上：从卡淡入的片段等淡化结束再出；后面接卡的片段在卡开始淡入前收掉
        lo = it["start"] + (it["xfn"] / FPS if k and "card" in items[k - 1] else 0)
        hi = items[k + 1]["start"] if k + 1 < len(items) and "card" in items[k + 1] else 1e9
        for a, b, txt in it.get("caps", []):
            caps.append([max(lo, it["start"] + (a - c0) / sp), min(hi, it["start"] + (b - c0) / sp), txt])
    caps.sort(key=lambda c: c[0])
    for i in range(len(caps) - 1):                   # 交叉淡化处相邻字幕不重叠
        caps[i][1] = min(caps[i][1], caps[i + 1][0] - 0.02)
    return items, caps, t


def cmd_plan(a, quiet=False):
    st = load_story(a.work)
    items, caps, total = plan(st)
    tgt = st.get("target")
    print(f"[i] 总时长 {total:.2f}s" + (f"（目标 {tgt}s）" if tgt else "") + f"，字幕 {len(caps)} 条")
    if tgt and abs(total - tgt) > 1:
        print(f"  [!] 与目标相差 {total - tgt:+.2f}s")
    if not quiet:
        for it in items:
            lab = it.get("card") or "{} {:.2f}-{:.2f} x{}".format(*it["clip"])
            print(f"    {it['start']:7.2f}s {it['n'] / FPS:5.2f}s xf={it['xfn'] / FPS:.1f}  {lab}")
    for c0, c1, txt in caps:
        from render import parse_marks
        txt = parse_marks(txt)[0]
        flag = ("  [!]<3s" if c1 - c0 < 2.95 else "") + ("  [!]>20字" if len(txt.replace(" ", "")) > 20 else "")
        if flag or not quiet:
            print(f"    {c0:7.2f}-{c1:7.2f}  {txt}{flag}")
    return st, items, caps, total


def cmd_cards(a):
    st = load_story(a.work)
    ck = st["_work"] / "checks"
    ck.mkdir(parents=True, exist_ok=True)
    for k, cfg in st["cards"].items():
        cd = Card(st, cfg)
        Image.fromarray(cd.frame(cd.n - 1)).save(ck / f"card_{k}.jpg", quality=90)
        print("[ok]", ck / f"card_{k}.jpg")


def make_caption(st, text):
    from render import parse_marks
    text, marked = parse_marks(text)
    f = font(st["fonts"]["caption"], int(40 * st["size"][0] / 1920))
    lh = int(f.size * 1.45)
    tw = int(f.getlength(text))
    w, h = tw + 64, lh + 28
    im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((0, 0, w - 1, h - 1), radius=18, fill=(12, 14, 20, int(255 * 0.74)))
    _, it, _, ib = f.getbbox("国")
    from render import draw_runs
    hl_c = tuple(st.get("caption_highlight_color", (255, 214, 102))) + (255,)
    hl = marked + list((st.get("caption_highlight_map") or {}).get(text, [])) + list(st.get("caption_highlight") or [])
    draw_runs(d, (w - tw) / 2, 14 + lh / 2 - (it + ib) / 2, text, f, hl, hl_fill=hl_c)
    return np.array(im)


def fmt(t):
    ms = int(round(t * 1000))
    return f"{ms // 3600000:02d}:{ms // 60000 % 60:02d}:{ms // 1000 % 60:02d},{ms % 1000:03d}"


def cmd_build(a):
    st, items, caps, total = cmd_plan(a, quiet=True)
    from render import parse_marks
    W, H = st["size"]
    name = st.get("name", "故事线")
    od = proj_path(st, st.get("out_dir", "output"))
    od.mkdir(parents=True, exist_ok=True)
    if a.out:
        out_sub = Path(a.out).expanduser()
    else:
        vs = [int(m.group(1)) for p in od.glob(f"{name}_v*.mp4") if (m := re.search(r"_v(\d+)\.mp4$", p.name))]
        out_sub = od / f"{name}_v{max(vs, default=0) + 1}.mp4"
    out_raw = out_sub.with_name(out_sub.stem + "_无字幕.mp4")

    def enc(path):
        return subprocess.Popen(["ffmpeg", "-y", "-v", "error", "-f", "rawvideo", "-pix_fmt", "rgb24", "-s", f"{W}x{H}",
                                 "-r", str(FPS), "-i", "-", "-c:v", "libx264", "-preset", "medium", "-crf", "17",
                                 "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(path)], stdin=subprocess.PIPE)

    e1, e2 = enc(out_sub), enc(out_raw)
    cap_imgs = [(c0, c1, make_caption(st, txt)) for c0, c1, txt in caps]
    gi = 0

    def emit(fr):
        nonlocal gi
        t = gi / FPS
        e2.stdin.write(np.ascontiguousarray(fr).tobytes())
        f2 = fr.copy()
        for c0, c1, img in cap_imgs:
            if c0 - 0.01 <= t <= c1:
                al = min(1.0, (t - c0) / 0.18 + 0.02, (c1 - t) / 0.18 + 0.02)
                blend_rgba(f2, img, (W - img.shape[1]) / 2, H - int(30 * W / 1920) - img.shape[0], max(0.0, al))
        e1.stdin.write(f2.tobytes())
        gi += 1
        if gi % (FPS * 10) == 0:
            print(f"  {gi / FPS:.0f}s / {total:.0f}s", flush=True)

    carry = []
    for k, it in enumerate(items):
        src = Card(st, st["cards"][it["card"]]).frames() if "card" in it else clip_frames(st, it["clip"])
        for j, prev in enumerate(carry):              # 与上一条交叉淡化
            al = (j + 1) / (len(carry) + 1)
            emit(cv2.addWeighted(prev, 1 - al, next(src), al, 0))
        nxt = items[k + 1]["xfn"] if k + 1 < len(items) else 0
        buf = deque()
        for idx, fr in enumerate(src, start=len(carry)):
            if idx >= it["n"] - nxt:
                buf.append(fr.copy())
            else:
                emit(fr)
        carry = list(buf)
    for fr in carry:
        emit(fr)
    for e in (e1, e2):
        e.stdin.close()
        e.wait()
    out_sub.with_suffix(".srt").write_text(
        "\n".join(f"{i}\n{fmt(c0)} --> {fmt(c1)}\n{parse_marks(txt)[0]}\n" for i, (c0, c1, txt) in enumerate(caps, 1)),
        encoding="utf-8")
    print(f"[ok] {out_sub}\n[ok] {out_raw}\n[ok] {out_sub.with_suffix('.srt')}  ({gi / FPS:.2f}s, {W}x{H})")


def main():
    ap = argparse.ArgumentParser()
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("ppt"); p.add_argument("pptx"); p.add_argument("--out", required=True)
    c = sp.add_parser("chapters"); c.add_argument("work"); c.add_argument("keys", nargs="*")
    for n in ("plan", "cards"):
        sp.add_parser(n).add_argument("work")
    b = sp.add_parser("build"); b.add_argument("work"); b.add_argument("--out")
    a = ap.parse_args()
    {"ppt": cmd_ppt, "chapters": cmd_chapters, "plan": cmd_plan, "cards": cmd_cards, "build": cmd_build}[a.cmd](a)


if __name__ == "__main__":
    main()
